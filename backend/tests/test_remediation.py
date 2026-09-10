import io
import pytest
import docx
from backend.services.document_service import extract_text, extract_text_from_docx
from backend.services.info_extractor import extract_education, extract_name
from backend.services.metadata_builder import build_analysis_metadata
from backend.services.scoring_service import rank_candidates
from backend.services.pipeline import AnalysisPipeline, AnalysisContext

def test_docx_table_extraction():
    """Verify that text within DOCX tables is extracted and not discarded."""
    doc = docx.Document()
    doc.add_paragraph("Alice Johnson")
    doc.add_paragraph("alice@example.com")
    
    # Add a table with education and skills
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Degree"
    table.cell(0, 1).text = "B.Tech in Computer Science"
    table.cell(1, 0).text = "Core Skills"
    table.cell(1, 1).text = "Python, FastAPI, Docker, PostgreSQL"
    table.cell(2, 0).text = "Experience"
    table.cell(2, 1).text = "5 years of backend engineering"
    
    bio = io.BytesIO()
    doc.save(bio)
    file_bytes = bio.getvalue()
    
    extracted = extract_text_from_docx(file_bytes)
    assert "Alice Johnson" in extracted
    assert "B.Tech in Computer Science" in extracted
    assert "Python, FastAPI, Docker, PostgreSQL" in extracted
    assert "5 years of backend engineering" in extracted

def test_doc_extension_rejection():
    """Verify that legacy .doc files raise a clean ValueError instead of crashing."""
    dummy_bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # OLE2 header magic bytes
    with pytest.raises(ValueError) as exc_info:
        extract_text(dummy_bytes, "resume.doc")
    assert "legacy binary format is not supported" in str(exc_info.value)

def test_metadata_builder_extracted_vs_matched_skills():
    """Verify metadata builder correctly separates all extracted skills from matched skills."""
    cand_dict = {
        "extracted_skills": ["python", "java", "docker", "react"],
        "matched_skills": ["python", "react"],
        "missing_skills": ["kubernetes"],
        "similarity_score": 85.0
    }
    meta = build_analysis_metadata(
        candidate_result=cand_dict,
        processing_time_ms=50,
        parser="docx",
        document_type="docx"
    )
    assert meta["skills"]["extracted"] == ["python", "java", "docker", "react"]
    assert meta["skills"]["matched"] == ["python", "react"]
    assert meta["skills"]["missing"] == ["kubernetes"]

@pytest.mark.anyio
async def test_full_pipeline_execution():
    """Verify that AnalysisPipeline executes all 7 stages and returns RecommendationBuiltEvent."""
    doc = docx.Document()
    doc.add_paragraph("Jane Smith")
    doc.add_paragraph("jane@example.com | 555-987-6543 | linkedin.com/in/janesmith")
    doc.add_paragraph("Education: M.S. in Software Engineering")
    doc.add_paragraph("Experience: 4 years building scalable services with Python, FastAPI, and Docker.")
    doc.add_paragraph("Projects: Built microservices platform with high availability.")
    bio = io.BytesIO()
    doc.save(bio)
    file_bytes = bio.getvalue()
    
    pipeline = AnalysisPipeline()
    jd = "Seeking a Software Engineer skilled in Python, FastAPI, and Docker."
    event = await pipeline.run_analysis(
        filename="jane_resume.docx",
        content_bytes=file_bytes,
        job_description=jd,
        clean_jd=jd
    )
    
    assert event.status == "success"
    assert event.explanation.scoring.candidate_name == "Jane Smith"
    assert event.explanation.scoring.candidate_email == "jane@example.com"
    assert event.explanation.scoring.candidate_education == "Master"
    assert event.explanation.scoring.similarity_score > 0
    assert len(event.recommendations) > 0

def test_scoring_service_reuse():
    """Verify rank_candidates reuses precalculated scores from ScoringStage without drift."""
    resumes = [
        {
            "filename": "cand1.docx",
            "name": "Alice",
            "similarity_score": 88.5,
            "skill_score": 90.0,
            "experience_score": 80.0,
            "education_score": 100.0,
            "projects_score": 80.0,
            "matched_skills": ["python", "docker"]
        },
        {
            "filename": "cand2.docx",
            "name": "Bob",
            "similarity_score": 65.0,
            "skill_score": 70.0,
            "experience_score": 50.0,
            "education_score": 60.0,
            "projects_score": 40.0,
            "matched_skills": ["python"]
        }
    ]
    ranked = rank_candidates(["python", "docker"], resumes)
    assert len(ranked) == 2
    assert ranked[0]["name"] == "Alice"
    assert ranked[0]["similarity_score"] == 88.5
    assert ranked[0]["rank"] == 1
    assert ranked[1]["name"] == "Bob"
    assert ranked[1]["similarity_score"] == 65.0
    assert ranked[1]["rank"] == 2
