import io
import PyPDF2
import docx
import logging

logger = logging.getLogger("resume_screener")

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file."""
    text = ""
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    except Exception as e:
        logger.error("Failed to read PDF: %s", e)
        raise ValueError(f"Failed to read PDF file format: {e}")
    logger.debug("Successfully extracted %d characters from PDF", len(text))
    return text.strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file, including paragraphs and tables."""
    text_parts = []
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate consecutive identical cells from merged columns
                deduped = []
                for cell_text in row_cells_text:
                    if not deduped or deduped[-1] != cell_text:
                        deduped.append(cell_text)
                if deduped:
                    text_parts.append(" | ".join(deduped))
    except Exception as e:
        logger.error("Failed to read DOCX: %s", e)
        raise ValueError(f"Failed to read DOCX file format: {e}")
    text = "\n".join(text_parts)
    logger.debug("Successfully extracted %d characters from DOCX", len(text))
    return text.strip()

def extract_text(file_bytes: bytes, filename: str) -> str:
    """Route to appropriate extractor based on extension."""
    filename_lower = filename.lower()
    if filename_lower.endswith('.pdf'):
        return extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith('.docx'):
        return extract_text_from_docx(file_bytes)
    elif filename_lower.endswith('.doc'):
        logger.warning("Unsupported legacy Word format rejected: %s", filename)
        raise ValueError("Unsupported file format: '.doc' legacy binary format is not supported. Please convert and upload your resume as '.docx' or '.pdf'.")
    else:
        # Fallback for plain text or unsupported formats
        try:
            return file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            raise ValueError("Unsupported file format or unreadable text encoding.")
